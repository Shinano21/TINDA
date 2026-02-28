from pos.models import Sales
from purchase.models import PurchaseProduct
from report.forms import *
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import uuid

from django.http import HttpResponse, HttpResponseBadRequest
from django.views.generic import FormView
from openpyxl import Workbook
from openpyxl.styles import Alignment
from django.db.models import Sum

from django.views import View  
class GenerateExcelProfitView(View):
    def post(self, request, *args, **kwargs):
        form = SalesReportForm(request.POST)
        if form.is_valid():
            sales_queryset = self.get_queryset(form)

            total_income = self.calculate_total_income(sales_queryset)
            total_costs = self.calculate_total_costs(sales_queryset)
            total_income_decimal = Decimal(total_income)
            total_costs_decimal = Decimal(total_costs)
            total_profit = total_income_decimal - total_costs_decimal

            sales_data, total_utilities = self.get_sales_data_and_utilities(sales_queryset)

            current_date = datetime.now()
            username = request.user.username
            unique_key = str(uuid.uuid4())

            
            excel_file = self.generate_excel_file(sales_data)

            
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
<<<<<<< HEAD
            response['Content-Disposition'] = f'attachment; filename="profit_report_general_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
=======
            response['Content-Disposition'] = f'attachment; filename="general_profit_report_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            response.write(excel_file.getvalue())

            return response
        else:
            return HttpResponseBadRequest("Invalid form")

    def get_queryset(self, form):
        queryset = Sales.objects.all()

        if form.is_valid():
            start_date = form.cleaned_data.get('start_date')
            end_date = form.cleaned_data.get('end_date')

            if start_date:
                queryset = queryset.filter(date_added__gte=start_date)
            if end_date:
                queryset = queryset.filter(date_added__lte=end_date)

        return queryset

    def calculate_total_income(self, sales_queryset):
        total_income = sales_queryset.aggregate(total=Sum('grand_total'))['total'] or Decimal('0')
        return total_income

    def calculate_total_costs(self, sales_queryset):
        total_costs = Decimal('0')

        for sale in sales_queryset:
            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    product_cost = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_cost = product_cost * Decimal(qty_bought)
                    total_costs += total_cost

        return total_costs

    def get_sales_data_and_utilities(self, sales_queryset):
        sales_data = []
        total_utilities = Decimal(0)

        for sale in sales_queryset:
            sale_cost = self.calculate_sale_cost(sale)
            sale_profit = Decimal(sale.grand_total) - sale_cost
            products_list = []

            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    cost_per_unit = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_qty_sold = item.qty
                    total_qty_bought = qty_bought

                    product_profit = (Decimal(item.qty) * sale_profit) / Decimal(total_qty_sold)

                    total_purchase_expense = cost_per_unit * Decimal(total_qty_bought)

                    gross_profit = (sale_cost + product_profit) - total_purchase_expense
                    total_utilities += gross_profit
                    products_list.append({
                        'product_name': item.product.name,
                        'cost_per_unit': cost_per_unit,
<<<<<<< HEAD
                        'total_qty_vendida': total_qty_vendida,
                        'total_qty_comprada': total_qty_comprada,
                        'product_ganancia': product_ganancia,
                        'ganancia_estado': 'Positive' if product_ganancia > 0 else ('Negative' if product_ganancia < 0 else 'Neutral'),
                        'total_gasto_compras': total_gasto_compras,
                        'ganancia_bruta': ganancia_bruta,
=======
                        'total_qty_sold': total_qty_sold,
                        'total_qty_bought': total_qty_bought,
                        'product_profit': product_profit,
                        'profit_status': 'Positive' if product_profit > 0 else ('Negative' if product_profit < 0 else 'Neutral'),
                        'total_purchase_expense': total_purchase_expense,
                        'gross_profit': gross_profit,
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
                    })

            sales_data.append({
                'date_added': sale.date_added,
                'products_list': products_list,
                'total_sale': Decimal(sale.grand_total),
                'total_cost': sale_cost,
                'total_profit': sale_profit,
            })

        return sales_data, total_utilities

    def calculate_sale_cost(self, sale):
        sale_cost = Decimal('0')
        for item in sale.salesitems_set.all():
            purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
            if purchase_product:
                product_cost = purchase_product.cost
                sale_cost += product_cost * Decimal(item.qty)
        return sale_cost

    def generate_excel_file(self, sales_data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Profit Report"

<<<<<<< HEAD
        # Headers
        headers = [
            "Sale Date", "Product Name", "Cost per Unit", "Qty Sold",
            "Qty Purchased", "Profit per Product", "Profit Status", "Total Purchase Cost",
=======
        # Header
        headers = [
            "Sale Date", "Product Name", "Cost per Unit", "Quantity Sold",
            "Quantity Bought", "Profit per Product", "Profit Status", "Total Purchase Expense",
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            "Gross Profit"
        ]
        ws.append(headers)

        # Data
        for sale in sales_data:
            for product in sale['products_list']:
                row = [
                    sale['date_added'].strftime('%Y-%m-%d %H:%M:%S'),
                    product['product_name'],
                    product['cost_per_unit'],
                    product['total_qty_sold'],
                    product['total_qty_bought'],
                    product['product_profit'],
                    product['profit_status'],
                    product['total_purchase_expense'],
                    product['gross_profit']
                ]
                ws.append(row)

        # Adjust column width
        for col in ws.iter_cols(min_col=1, max_col=ws.max_column):
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column].width = adjusted_width


        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        return excel_file


class YearlyExcelProfitView(View):
    def post(self, request, *args, **kwargs):
        form = YearReportForm(request.POST)
        if form.is_valid():
            year = form.cleaned_data.get('year')
            sales_queryset = self.get_queryset(year)

            total_income = self.calculate_total_income(sales_queryset)
            total_costs = self.calculate_total_costs(sales_queryset)
            total_income_decimal = Decimal(total_income)
            total_costs_decimal = Decimal(total_costs)
            total_profit = total_income_decimal - total_costs_decimal

            sales_data, total_utilities = self.get_sales_data_and_utilities(sales_queryset)

            current_date = datetime.now()
            username = request.user.username
            unique_key = str(uuid.uuid4())


            excel_file = self.generate_excel_file(sales_data)


            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
<<<<<<< HEAD
            response['Content-Disposition'] = f'attachment; filename="profit_report_yearly_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
=======
            response['Content-Disposition'] = f'attachment; filename="yearly_profit_report_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            response.write(excel_file.getvalue())

            return response
        else:
            return HttpResponseBadRequest("Invalid form")

    def get_queryset(self, year):
        return Sales.objects.filter(date_added__year=year)

    def calculate_total_income(self, sales_queryset):
        total_income = sales_queryset.aggregate(total=Sum('grand_total'))['total'] or Decimal('0')
        return total_income

    def calculate_total_costs(self, sales_queryset):
        total_costs = Decimal('0')

        for sale in sales_queryset:
            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    product_cost = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_cost = product_cost * Decimal(qty_bought)
                    total_costs += total_cost

        return total_costs

    def get_sales_data_and_utilities(self, sales_queryset):
        sales_data = []
        total_utilities = Decimal(0)

        for sale in sales_queryset:
            sale_cost = self.calculate_sale_cost(sale)
            sale_profit = Decimal(sale.grand_total) - sale_cost
            products_list = []

            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    cost_per_unit = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_qty_sold = item.qty
                    total_qty_bought = qty_bought

                    product_profit = (Decimal(item.qty) * sale_profit) / Decimal(total_qty_sold)

                    total_purchase_expense = cost_per_unit * Decimal(total_qty_bought)

                    gross_profit = (sale_cost + product_profit) - total_purchase_expense
                    total_utilities += gross_profit
                    products_list.append({
                        'product_name': item.product.name,
                        'cost_per_unit': cost_per_unit,
<<<<<<< HEAD
                        'total_qty_vendida': total_qty_vendida,
                        'total_qty_comprada': total_qty_comprada,
                        'product_ganancia': product_ganancia,
                        'ganancia_estado': 'Positive' if product_ganancia > 0 else ('Negative' if product_ganancia < 0 else 'Neutral'),
                        'total_gasto_compras': total_gasto_compras,
                        'ganancia_bruta': ganancia_bruta,
=======
                        'total_qty_sold': total_qty_sold,
                        'total_qty_bought': total_qty_bought,
                        'product_profit': product_profit,
                        'profit_status': 'Positive' if product_profit > 0 else ('Negative' if product_profit < 0 else 'Neutral'),
                        'total_purchase_expense': total_purchase_expense,
                        'gross_profit': gross_profit,
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
                    })

            sales_data.append({
                'date_added': sale.date_added,
                'products_list': products_list,
                'total_sale': Decimal(sale.grand_total),
                'total_cost': sale_cost,
                'total_profit': sale_profit,
            })

        return sales_data, total_utilities

    def calculate_sale_cost(self, sale):
        sale_cost = Decimal('0')
        for item in sale.salesitems_set.all():
            purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
            if purchase_product:
                product_cost = purchase_product.cost
                sale_cost += product_cost * Decimal(item.qty)
        return sale_cost

    def generate_excel_file(self, sales_data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Profit Report"

    
        headers = [
            "Sale Date", "Product Name", "Cost per Unit", "Quantity Sold",
            "Quantity Bought", "Profit per Product", "Profit Status", "Total Purchase Expense",
            "Gross Profit"
        ]
        ws.append(headers)

    
        for sale in sales_data:
            for product in sale['products_list']:
                row = [
                    sale['date_added'].strftime('%Y-%m-%d %H:%M:%S'),
                    product['product_name'],
                    product['cost_per_unit'],
                    product['total_qty_sold'],
                    product['total_qty_bought'],
                    product['product_profit'],
                    product['profit_status'],
                    product['total_purchase_expense'],
                    product['gross_profit']
                ]
                ws.append(row)

    
        for col in ws.iter_cols(min_col=1, max_col=ws.max_column):
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column].width = adjusted_width

    
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        return excel_file


class MonthlyExcelProfitView(FormView):
    def post(self, request, *args, **kwargs):
        form = MonthYearReportForm(request.POST)
        if form.is_valid():
            year = form.cleaned_data.get('year')
            month = form.cleaned_data.get('month')
            sales_queryset = self.get_queryset(year, month)

            total_income = self.calculate_total_income(sales_queryset)
            total_costs = self.calculate_total_costs(sales_queryset)
            total_income_decimal = Decimal(total_income)
            total_costs_decimal = Decimal(total_costs)
            total_profit = total_income_decimal - total_costs_decimal

            sales_data, total_utilities = self.get_sales_data_and_utilities(sales_queryset)

            current_date = datetime.now()
            username = request.user.username
            unique_key = str(uuid.uuid4())

        
            excel_file = self.generate_excel_file(sales_data)

        
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
<<<<<<< HEAD
            response['Content-Disposition'] = f'attachment; filename="profit_report_monthly_{month}_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
=======
            response['Content-Disposition'] = f'attachment; filename="monthly_profit_report_{month}_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            response.write(excel_file.getvalue())

            return response
        else:
            return HttpResponseBadRequest("Invalid form")

    def get_queryset(self, year, month):
        return Sales.objects.filter(date_added__year=year, date_added__month=month)

    def calculate_total_income(self, sales_queryset):
        total_income = sales_queryset.aggregate(total=Sum('grand_total'))['total'] or Decimal('0')
        return total_income

    def calculate_total_costs(self, sales_queryset):
        total_costs = Decimal('0')

        for sale in sales_queryset:
            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    product_cost = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_cost = product_cost * Decimal(qty_bought)
                    total_costs += total_cost

        return total_costs

    def get_sales_data_and_utilities(self, sales_queryset):
        sales_data = []
        total_utilities = Decimal(0)

        for sale in sales_queryset:
            sale_cost = self.calculate_sale_cost(sale)
            sale_profit = Decimal(sale.grand_total) - sale_cost
            products_list = []

            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    cost_per_unit = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_qty_sold = item.qty
                    total_qty_bought = qty_bought

                    product_profit = (Decimal(item.qty) * sale_profit) / Decimal(total_qty_sold)

                    total_purchase_expense = cost_per_unit * Decimal(total_qty_bought)

                    gross_profit = (sale_cost + product_profit) - total_purchase_expense
                    total_utilities += gross_profit
                    products_list.append({
                        'product_name': item.product.name,
                        'cost_per_unit': cost_per_unit,
                        'total_qty_sold': total_qty_sold,
                        'total_qty_bought': total_qty_bought,
                        'product_profit': product_profit,
                        'profit_status': 'Positive' if product_profit > 0 else ('Negative' if product_profit < 0 else 'Neutral'),
                        'total_purchase_expense': total_purchase_expense,
                        'gross_profit': gross_profit,
                    })

            sales_data.append({
                'date_added': sale.date_added,
                'products_list': products_list,
                'total_sale': Decimal(sale.grand_total),
                'total_cost': sale_cost,
                'total_profit': sale_profit,
            })

        return sales_data, total_utilities

    def calculate_sale_cost(self, sale):
        sale_cost = Decimal('0')
        for item in sale.salesitems_set.all():
            purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
            if purchase_product:
                product_cost = purchase_product.cost
                sale_cost += product_cost * Decimal(item.qty)
        return sale_cost

    def generate_excel_file(self, sales_data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Profit Report"

        headers = [
<<<<<<< HEAD
            "Sale Date", "Product Name", "Cost per Unit", "Qty Sold",
            "Qty Purchased", "Profit per Product", "Profit Status", "Total Purchase Cost",
=======
            "Sale Date", "Product Name", "Cost per Unit", "Quantity Sold",
            "Quantity Bought", "Profit per Product", "Profit Status", "Total Purchase Expense",
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            "Gross Profit"
        ]
        ws.append(headers)

        
        for sale in sales_data:
            for product in sale['products_list']:
                row = [
                    sale['date_added'].strftime('%Y-%m-%d %H:%M:%S'),
                    product['product_name'],
                    product['cost_per_unit'],
                    product['total_qty_sold'],
                    product['total_qty_bought'],
                    product['product_profit'],
                    product['profit_status'],
                    product['total_purchase_expense'],
                    product['gross_profit']
                ]
                ws.append(row)


        for col in ws.iter_cols(min_col=1, max_col=ws.max_column):
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column].width = adjusted_width


        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        return excel_file


class DailyExcelProfitView(FormView):
    template_name = 'your_template.html'
    form_class = DayMonthYearReportForm

    def form_valid(self, form):
        year = form.cleaned_data.get('year')
        month = form.cleaned_data.get('month')
        day = form.cleaned_data.get('day')

        sales_queryset = self.get_queryset(year, month, day)

        total_income = self.calculate_total_income(sales_queryset)
        total_costs = self.calculate_total_costs(sales_queryset)
        total_income_decimal = Decimal(total_income)
        total_costs_decimal = Decimal(total_costs)
        total_profit = total_income_decimal - total_costs_decimal

        sales_data, total_utilities = self.get_sales_data_and_utilities(sales_queryset)

        current_date = datetime.now()
        username = self.request.user.username
        unique_key = str(uuid.uuid4())

        
        excel_file = self.generate_excel_file(sales_data)

        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
<<<<<<< HEAD
        response['Content-Disposition'] = f'attachment; filename="profit_report_daily_{day}_{month}_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
=======
        response['Content-Disposition'] = f'attachment; filename="daily_profit_report_{day}_{month}_{year}_{current_date.strftime("%Y%m%d_%H%M%S")}.xlsx"'
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
        response.write(excel_file.getvalue())

        return response

    def get_queryset(self, year, month, day):
        queryset = Sales.objects.filter(date_added__year=year, date_added__month=month)

        if day:
            queryset = queryset.filter(date_added__day=day)

        return queryset

    def calculate_total_income(self, sales_queryset):
        total_income = sales_queryset.aggregate(total=Sum('grand_total'))['total'] or Decimal('0')
        return total_income

    def calculate_total_costs(self, sales_queryset):
        total_costs = Decimal('0')

        for sale in sales_queryset:
            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    product_cost = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_cost = product_cost * Decimal(qty_bought)
                    total_costs += total_cost

        return total_costs

    def get_sales_data_and_utilities(self, sales_queryset):
        sales_data = []
        total_utilities = Decimal(0)

        for sale in sales_queryset:
            sale_cost = self.calculate_sale_cost(sale)
            sale_profit = Decimal(sale.grand_total) - sale_cost
            products_list = []

            for item in sale.salesitems_set.all():
                purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
                if purchase_product:
                    cost_per_unit = purchase_product.cost
                    qty_bought = sum([pp.qty for pp in PurchaseProduct.objects.filter(product=item.product)])
                    total_qty_sold = item.qty
                    total_qty_bought = qty_bought

                    product_profit = (Decimal(item.qty) * sale_profit) / Decimal(total_qty_sold)

                    total_purchase_expense = cost_per_unit * Decimal(total_qty_bought)

                    gross_profit = (sale_cost + product_profit) - total_purchase_expense
                    total_utilities += gross_profit
                    products_list.append({
                        'product_name': item.product.name,
                        'cost_per_unit': cost_per_unit,
<<<<<<< HEAD
                        'total_qty_vendida': total_qty_vendida,
                        'total_qty_comprada': total_qty_comprada,
                        'product_ganancia': product_ganancia,
                        'ganancia_estado': 'Positive' if product_ganancia > 0 else ('Negative' if product_ganancia < 0 else 'Neutral'),
                        'total_gasto_compras': total_gasto_compras,
                        'ganancia_bruta': ganancia_bruta,
=======
                        'total_qty_sold': total_qty_sold,
                        'total_qty_bought': total_qty_bought,
                        'product_profit': product_profit,
                        'profit_status': 'Positive' if product_profit > 0 else ('Negative' if product_profit < 0 else 'Neutral'),
                        'total_purchase_expense': total_purchase_expense,
                        'gross_profit': gross_profit,
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
                    })

            sales_data.append({
                'date_added': sale.date_added,
                'products_list': products_list,
                'total_sale': Decimal(sale.grand_total),
                'total_cost': sale_cost,
                'total_profit': sale_profit,
            })

        return sales_data, total_utilities

    def calculate_sale_cost(self, sale):
        sale_cost = Decimal('0')
        for item in sale.salesitems_set.all():
            purchase_product = PurchaseProduct.objects.filter(product=item.product).first()
            if purchase_product:
                product_cost = purchase_product.cost
                sale_cost += product_cost * Decimal(item.qty)
        return sale_cost

    def generate_excel_file(self, sales_data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Profit Report"

        headers = [
<<<<<<< HEAD
            "Sale Date", "Product Name", "Cost per Unit", "Qty Sold",
            "Qty Purchased", "Profit per Product", "Profit Status", "Total Purchase Cost",
=======
            "Sale Date", "Product Name", "Cost per Unit", "Quantity Sold",
            "Quantity Bought", "Profit per Product", "Profit Status", "Total Purchase Expense",
>>>>>>> 34cc06a43306d90c1ec90c9b7ceefa3e8a99c7ea
            "Gross Profit"
        ]
        ws.append(headers)

        
        for sale in sales_data:
            for product in sale['products_list']:
                row = [
                    sale['date_added'].strftime('%Y-%m-%d %H:%M:%S'),
                    product['product_name'],
                    product['cost_per_unit'],
                    product['total_qty_sold'],
                    product['total_qty_bought'],
                    product['product_profit'],
                    product['profit_status'],
                    product['total_purchase_expense'],
                    product['gross_profit']
                ]
                ws.append(row)

        
        for col in ws.iter_cols(min_col=1, max_col=ws.max_column):
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column].width = adjusted_width

        
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        return excel_file
